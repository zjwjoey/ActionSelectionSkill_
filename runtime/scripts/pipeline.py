"""Offline, name-only pilot. Production inputs are opened read-only."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import random
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
import storage

ROOT = Path(__file__).resolve().parents[1]
RULE_VERSION = "name-semantics-v2"
FIELDS = ("category_l1", "category_l2", "product_type", "primary_function",
          "usage_scenario", "consumer_need", "action_pattern", "confidence", "evidence", "review_status")
MULTI = ("primary_function", "usage_scenario", "consumer_need", "action_pattern")
EXTRA_FIELDS = ("uncertainty_codes", "dimension_confidence", "ai_review")
STATUSES = {"待审核", "已确认", "需修改", "信息不足"}
STRATA = dict(zip(
    ["家居", "个人护理", "时尚", "厨房", "兴趣爱好", "家居维修", "食品饮料", "玩具", "家庭日用", "办公文具", "园艺", "数码多媒体", "宠物", "运动用品", "旅行"],
    ["Vivienda", "Cuidado personal", "Moda", "Cocina", "Hobby", "Bricolaje", "Comer y beber", "Juguetes", "Hogar", "Oficina y papelería", "Jardín", "Multimedia", "Mascotas", "Artículos deportivos", "Viajes"]))


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def digest(value):
    return hashlib.sha256(canonical(value).encode("utf-8")).hexdigest()


def filehash(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def read_jsonl(path):
    with Path(path).open(encoding="utf-8-sig") as f:
        return [json.loads(line) for line in f if line.strip()]


def immutable(path, value, *, jsonl=False):
    """Create once; same-content retries succeed, differing retries fail."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    body = ("".join(canonical(x) + "\n" for x in value) if jsonl else json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n")
    data = body.encode("utf-8")
    if path.exists():
        if path.read_bytes() != data:
            raise ValueError(f"IMMUTABLE_ARTIFACT_CONFLICT: {path}")
        return
    tmp = path.with_name(path.name + ".partial")
    with tmp.open("wb") as f:
        f.write(data)
        f.flush()
        os.fsync(f.fileno())
    # Hard link publishes without replacing an existing immutable artifact.
    try:
        os.link(tmp, path)
    except FileExistsError:
        if path.read_bytes() != data:
            raise ValueError(f"IMMUTABLE_ARTIFACT_CONFLICT: {path}")
    finally:
        tmp.unlink(missing_ok=True)


def normalize_name(value):
    # Preserve accents, brands, numbers, interfaces, negation and qualifiers.
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(value or ""))).strip()


def sku_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def guarded_output(path):
    if not Path(path).resolve().is_relative_to(ROOT.resolve()):
        raise ValueError("OUTPUT_MUST_STAY_IN_INDEPENDENT_PROJECT")
    return Path(path)


def snapshot(master):
    from openpyxl import load_workbook
    master = Path(master).resolve()
    if master.is_relative_to(ROOT):
        raise ValueError("SOURCE_MUST_BE_EXTERNAL")
    before = {"sha256": filehash(master), "bytes": master.stat().st_size, "mtime_ns": master.stat().st_mtime_ns}
    sid = "master_" + before["sha256"][:16]
    out = ROOT / "source_snapshot" / sid
    if (out / "manifest.json").exists():
        validate_snapshot(sid)
        result = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
        storage.register_snapshot(ROOT, result, read_jsonl(out / "sku_source_rows.jsonl"), read_jsonl(out / "archive_reference_rows.jsonl"), read_jsonl(ROOT / "learning_data" / sid / "base.jsonl"))
        return result
    rows, archive = [], []
    workbook = load_workbook(master, read_only=True, data_only=True)
    try:
        for sn, header_row, fields in [
            ("02_SKU_ES_CURRENT", 1, {"sku": "SKU", "name": "西班牙语品名", "category": "一级类目（西语）", "cat2": "二级类目（西语）", "price": "当前售价 (€)", "date": "最后确认存在日期", "reference": "Canonical_ID"}),
            ("08_LONG_TERM_MASTER", 7, {"sku": "正式SKU", "name": "西班牙语品名", "category": "一级类目（西语）", "price": "当前售价 (€)", "date": "最后观察日期", "reference": "来源工作表"}),
            ("07_APRIL_ARCHIVE", 1, {"sku": "正式SKU", "name": "西班牙语品名", "category": "一级类目（西语）", "price": "历史售价 (€)", "date": "观测日期", "reference": "西语来源文件"}),
        ]:
            sheet = workbook[sn]
            it = sheet.iter_rows(values_only=True)
            for _ in range(header_row):
                header = next(it)
            idx = {v: i for i, v in enumerate(header) if v}
            for required in fields.values():
                if required not in idx:
                    raise ValueError(f"UNKNOWN_MASTER_SCHEMA: {sn}/{required}")
            for rownum, row in enumerate(it, header_row + 1):
                if not any(v is not None for v in row):
                    continue
                get = lambda k: row[idx[fields[k]]] if k in fields else None
                sku = sku_text(get("sku"))
                source_reference = {"master": str(master), "master_sha256": before["sha256"], "sheet": sn, "excel_row": rownum, "original_reference": str(get("reference") or "")}
                entry = {"source_row_id": "row_" + digest([before["sha256"], sn, rownum])[:24], "sku": sku or None,
                         "name_es_original": get("name"), "category_original": get("category"), "category_l2_original": get("cat2"),
                         "price_original": get("price"), "source_date": str(get("date") or ""), "source_reference": source_reference,
                         "snapshot_id": sid, "raw_row_sha256": digest(list(row))}
                if sn == "07_APRIL_ARCHIVE" or not sku:
                    entry["exclusion_reason"] = "ARCHIVE_REFERENCE_ONLY_NO_CONFIRMED_SKU" if not sku else "HISTORICAL_ARCHIVE_REFERENCE"
                    entry["archive_identifier"] = str(row[0] or "")
                    archive.append(entry)
                else:
                    rows.append(entry)
    finally:
        workbook.close()
    after = {"sha256": filehash(master), "bytes": master.stat().st_size, "mtime_ns": master.stat().st_mtime_ns}
    if before != after:
        raise ValueError("SOURCE_CHANGED_DURING_READ: no snapshot published")
    grouped = defaultdict(list)
    for row in rows:
        grouped[row["sku"]].append(row)
    learning = []
    for sku, versions in sorted(grouped.items()):
        candidates = sorted(versions, key=lambda r: (0 if r["source_reference"]["sheet"] == "02_SKU_ES_CURRENT" else 1, r["source_row_id"]))
        chosen = next((r for r in candidates if normalize_name(r["name_es_original"])), candidates[0])
        item = dict(chosen)
        item["source_row_ids"] = [r["source_row_id"] for r in versions]
        item["source_name_variant_count"] = len({str(r["name_es_original"] or "") for r in versions})
        item["name_es_normalized"] = normalize_name(chosen["name_es_original"])
        item["analysis_status"] = "未处理"
        learning.append(item)
    assert len(rows) == len({r["source_row_id"] for r in rows})
    immutable(out / "sku_source_rows.jsonl", rows, jsonl=True)
    immutable(out / "archive_reference_rows.jsonl", archive, jsonl=True)
    work = ROOT / "learning_data" / sid / "base.jsonl"
    immutable(work, learning, jsonl=True)
    manifest = {"snapshot_id": sid, "source": str(master), "source_before": before, "source_after": after,
                "source_unchanged": True, "authority": "AGENTS.md identifies this as formal working Master; settings.yaml declares SQLITE_PRIMARY. This snapshot is the workbook projection, not a copy of all database revision history.",
                "sku_source_rows": len(rows), "unique_skus": len(grouped), "source_rows_by_sheet": dict(Counter(r["source_reference"]["sheet"] for r in rows)),
                "archive_reference_rows": len(archive), "archive_rows_without_confirmed_sku": sum(not r["sku"] for r in archive),
                "missing_names_source_rows": sum(not normalize_name(r["name_es_original"]) for r in rows),
                "nonempty_learning_names": sum(bool(r["name_es_normalized"]) for r in learning),
                "multiple_source_rows_skus": sum(len(v) > 1 for v in grouped.values()),
                "different_name_versions_skus": sum(x["source_name_variant_count"] > 1 for x in learning),
                "same_name_different_skus": sum(n > 1 for n in Counter(x["name_es_normalized"].casefold() for x in learning if x["name_es_normalized"]).values()),
                "scope": "All confirmed SKU rows in 02/08 retained, including repeated SKU sources; 07 and unconfirmed 08 entities retained separately as archive references. Prices stored but never fed to semantic analysis. Images, descriptions, details, sales, stock and database event history excluded.",
                "files": {str(p.relative_to(ROOT)): filehash(p) for p in [out / "sku_source_rows.jsonl", out / "archive_reference_rows.jsonl", work]}}
    immutable(out / "manifest.json", manifest)
    storage.register_snapshot(ROOT, manifest, rows, archive, learning)
    validate_snapshot(sid)
    immutable(ROOT / "reports" / sid / "p0_integrity.json", {**manifest, "reload_and_mapping_check": "PASS"})
    return manifest


def validate_snapshot(sid):
    manifest = json.loads((ROOT / "source_snapshot" / sid / "manifest.json").read_text(encoding="utf-8"))
    for name, expected in manifest["files"].items():
        if filehash(ROOT / name) != expected:
            raise ValueError("SNAPSHOT_HASH_MISMATCH: " + name)
    rows = read_jsonl(ROOT / "source_snapshot" / sid / "sku_source_rows.jsonl")
    base = read_jsonl(ROOT / "learning_data" / sid / "base.jsonl")
    by_id = {x["source_row_id"]: x for x in rows}
    assert len(by_id) == len(rows) == manifest["sku_source_rows"]
    assert {x["sku"] for x in rows} == {x["sku"] for x in base}
    assert len(base) == len({x["sku"] for x in base}) == manifest["unique_skus"]
    assert {rid for x in base for rid in x["source_row_ids"]} == set(by_id)
    for x in base:
        assert all(by_id[rid]["sku"] == x["sku"] for rid in x["source_row_ids"])
        assert x["name_es_original"] == by_id[x["source_row_id"]]["name_es_original"]
    return manifest


def sample(sid, seed=20261022, count=100):
    if count != 100:
        raise ValueError("P1_SCOPE_IS_EXACTLY_100; expand only after Owner review")
    validate_snapshot(sid)
    rows = [x for x in storage.load_products(ROOT, sid) if x["name_es_normalized"]]
    groups = defaultdict(list)
    for x in rows:
        key = str(x["category_original"] or "未分类")
        groups[STRATA.get(key, key)].append(x)
    if len(groups) > count:
        raise ValueError("STRATA_EXCEED_SAMPLE_SIZE")
    quotas = {k: min(6, len(v)) for k, v in groups.items()}
    while sum(quotas.values()) < count:
        k = max((k for k in groups if quotas[k] < len(groups[k])), key=lambda k: math.sqrt(len(groups[k])) / (quotas[k] + 1))
        quotas[k] += 1
    selected = []
    for category, pool in sorted(groups.items()):
        rng = random.Random(str(seed) + category)
        pool = sorted(pool, key=lambda x: x["sku"])
        rng.shuffle(pool)
        head = lambda x: " ".join(x["name_es_normalized"].casefold().split()[:2])
        freq = Counter(head(x) for x in pool)
        candidates = [(min(pool, key=lambda x: len(x["name_es_normalized"])), "该来源分层中最短品名，覆盖模糊名称"),
                      (max(pool, key=lambda x: len(x["name_es_normalized"])), "该来源分层中最长品名，覆盖复合用途"),
                      (min(pool, key=lambda x: freq[head(x)]), "稀少名称词头，覆盖长尾类型"),
                      (max(pool, key=lambda x: freq[head(x)]), "常见名称词头，覆盖常见类型")]
        candidates += [(x, "固定种子分层随机补足") for x in pool]
        seen = set()
        for x, why in candidates:
            if x["sku"] in seen:
                continue
            seen.add(x["sku"])
            selected.append({**x, "sampling_stratum": category, "sampling_reason": why, "sampling_seed": seed})
            if len(seen) == quotas[category]:
                break
    assert len(selected) == count == len({x["sku"] for x in selected})
    sample_id = "pilot100_" + digest([sid, seed, [x["sku"] for x in selected]])[:12]
    folder = ROOT / "learning_data" / sid / sample_id
    immutable(folder / "sample_provenance.jsonl", selected, jsonl=True)
    # Semantic analyst receives ONLY identity and name. No original category, price, date, status, or Chinese.
    blind = [{"sku": x["sku"], "name_es_original": x["name_es_original"], "name_es_normalized": x["name_es_normalized"], "source_row_id": x["source_row_id"], "snapshot_id": sid} for x in selected]
    immutable(folder / "names_only_input.jsonl", blind, jsonl=True)
    manifest = {"sample_id": sample_id, "snapshot_id": sid, "seed": seed, "count": count, "quotas": quotas,
                "category_use": "Sampling and post-hoc comparison only. Hidden from semantic analyst.",
                "selection": "Per stratum: shortest, longest, rare/common name head, seeded random; sqrt allocation for remaining slots.",
                "input_sha256": filehash(folder / "names_only_input.jsonl")}
    immutable(folder / "sample_manifest.json", manifest)
    storage.register_batch(ROOT,manifest,selected,blind)
    return manifest


def get_folder(sid, batch):
    folder = guarded_output(ROOT / "learning_data" / sid / batch)
    manifest = json.loads((folder / "sample_manifest.json").read_text(encoding="utf-8"))
    if manifest["count"] != 100 or filehash(folder / "names_only_input.jsonl") != manifest["input_sha256"]:
        raise ValueError("INVALID_PILOT_MANIFEST")
    validate_snapshot(sid)
    return folder


def state(folder):
    return storage.state(ROOT,folder.name)


def validate_annotation(x, original, rules_version=None):
    for key in FIELDS:
        if key not in x:
            raise ValueError("MISSING_ANNOTATION_FIELD: " + key)
    if x["review_status"] not in {"待审核", "信息不足"}:
        raise ValueError("AI_CANNOT_CONFIRM_ITSELF")
    if x["confidence"] not in {"高", "中", "低"}:
        raise ValueError("INVALID_CONFIDENCE")
    for f in MULTI:
        if not isinstance(x[f], list) or any(not isinstance(v, str) or not v.strip() for v in x[f]):
            raise ValueError("INVALID_MULTILABEL: " + f)
    for f in ("category_l1", "category_l2", "product_type"):
        if not isinstance(x[f], str) or not x[f].strip():
            raise ValueError("INVALID_CATEGORY_OR_TYPE: " + f)
    ev = x["evidence"]
    if not isinstance(ev, dict) or not ev.get("title_spans") or not ev.get("inference_boundary"):
        raise ValueError("EVIDENCE_MUST_BIND_TO_TITLE")
    title = normalize_name(original["name_es_original"]).casefold()
    if not isinstance(ev["title_spans"],list) or any(not isinstance(v,str) or not normalize_name(v) or normalize_name(v).casefold() not in title for v in ev["title_spans"]):
        raise ValueError("EVIDENCE_SPAN_NOT_IN_OWN_NAME")
    if x["review_status"] == "信息不足" and x["confidence"] != "低":
        raise ValueError("INSUFFICIENT_INFORMATION_REQUIRES_LOW_CONFIDENCE")
    version=rules_version or x.get('rules_version',RULE_VERSION)
    if version == 'name-semantics-v2':
        codes=x.get('uncertainty_codes');dims=x.get('dimension_confidence')
        definitions=load_rules(version)
        nr=definitions['name_rules.json']
        if not isinstance(codes,list) or any(c not in nr['uncertainty_codes'] for c in codes):
            raise ValueError('INVALID_UNCERTAINTY_CODES')
        if not isinstance(dims,dict) or set(dims)!=set(nr['dimension_confidence_fields']) or any(v not in {'高','中','低'} for v in dims.values()):
            raise ValueError('INVALID_DIMENSION_CONFIDENCE')
        blocks=set(codes)&set(nr['blocking_uncertainty_codes'])
        if bool(blocks)!=(x['review_status']=='信息不足'):
            raise ValueError('BLOCKING_UNCERTAINTY_STATUS_MISMATCH')
        if re.search(r'[\u3400-\u9fff]',title) and not re.search(r'[a-zà-ÿ]',title) and 'SOURCE_LANGUAGE_INVALID' not in codes:
            raise ValueError('SOURCE_LANGUAGE_INVALID_NOT_RECORDED')
        if x['category_l1'] not in definitions['taxonomy.json']['category_l1_candidates']:
            raise ValueError('UNKNOWN_CATEGORY_L1')
        if any(t not in definitions['function_tags.json']['action_pattern_candidates'] for t in x['action_pattern']):
            raise ValueError('UNKNOWN_ACTION_PATTERN')
        review=x.get('ai_review')
        if review and (review.get('reviewer_kind')!='AI' or review.get('decision') not in definitions['taxonomy.json']['review_decisions'] or not review.get('reason')):
            raise ValueError('INVALID_AI_REVIEW_METADATA')


def load_rules(version):
    folder=ROOT/'rules'/'versions'/version
    if not folder.exists():folder=ROOT/'rules'
    definitions={p.name:json.loads(p.read_text(encoding='utf-8')) for p in sorted(folder.glob('*.json'))}
    if not definitions or any(x.get('version')!=version for x in definitions.values()):
        raise ValueError('RULE_VERSION_CONTENT_MISMATCH')
    return definitions


def import_analysis(sid, batch, path, reprocess=(), rules_version=None):
    folder = get_folder(sid, batch)
    blind = {x["sku"]: x for x in storage.members(ROOT,batch,blind=True)}
    version=rules_version or RULE_VERSION
    definitions=load_rules(version)
    storage.register_rule(ROOT,version,digest(definitions),definitions)
    existing, feedback = state(folder)
    incoming = read_jsonl(path)
    if len({x["sku"] for x in incoming}) != len(incoming):
        raise ValueError("DUPLICATE_INCOMING_SKU")
    allowed = set(reprocess)
    clean = []
    for x in incoming:
        sku = x["sku"]
        if sku not in blind:
            raise ValueError("SKU_OUTSIDE_PILOT: " + sku)
        validate_annotation(x, blind[sku],version)
        payload = {k: x[k] for k in (*FIELDS,*EXTRA_FIELDS) if k in x}
        # New versioned rule contracts can state a more precise generation method.
        # Preserve the existing label for historical versions that lack this key.
        method = definitions["name_rules.json"].get("analysis_method_label", "Codex title-only semantic analysis; no parameter training or external API")
        proposal = {**blind[sku], **payload, "rules_version": version, "analysis_method": method, "annotation_hash": digest(payload), "human_confirmed": False}
        if sku in existing:
            if existing[sku]["annotation_hash"] == proposal["annotation_hash"] and existing[sku]["rules_version"] == version:
                continue
            if sku not in allowed:
                raise ValueError("EXPLICIT_REPROCESS_REQUIRED: " + sku)
            if feedback.get(sku, {}).get("review_status") == "已确认":
                raise ValueError("CONFIRMED_SKU_PROTECTED: " + sku)
        clean.append(proposal)
    if clean:
        immutable(folder / "analysis" / ("analysis_" + digest(clean)[:20] + ".jsonl"), clean, jsonl=True)
        storage.put_analyses(ROOT,batch,clean)
    return {"imported": len(clean), "already_present": len(incoming) - len(clean), "total_current": len(state(folder)[0])}


def import_feedback(sid, batch, path):
    folder = get_folder(sid, batch)
    proposals, previous = state(folder)
    if Path(path).suffix.lower() == ".csv":
        rows = []
        with Path(path).open(encoding="utf-8-sig",newline="") as f:
            for r in csv.DictReader(f):
                reviewer=str(r.get("human_reviewer","") or "").strip()
                status=str(r.get("human_review_status","") or "").strip()
                if not reviewer and status in {"", "待审核"}:
                    continue
                rows.append({"sku":r["sku"],"expected_annotation_hash":r["annotation_hash"],"review_status":status,
                             "corrections":json.loads(r.get("human_corrections") or "{}"),"reviewer":reviewer,"note":r.get("human_review_note","")})
    else:
        rows = read_jsonl(path)
    if len({x["sku"] for x in rows}) != len(rows):
        raise ValueError("DUPLICATE_FEEDBACK_SKU")
    clean = []
    for x in rows:
        if x["sku"] not in proposals or x.get("expected_annotation_hash") != proposals[x["sku"]]["annotation_hash"]:
            raise ValueError("STALE_OR_UNKNOWN_FEEDBACK")
        if not str(x.get("reviewer", "")).strip() or x.get("review_status") not in STATUSES:
            raise ValueError("FEEDBACK_REVIEWER_AND_STATUS_REQUIRED")
        if any(k not in FIELDS for k in x.get("corrections", {})):
            raise ValueError("FEEDBACK_CANNOT_CHANGE_SOURCE_FIELDS")
        x = {**x, "feedback_contract": "explicit-human-feedback-v1"}
        if x != previous.get(x["sku"]):
            clean.append(x)
    if clean:
        immutable(folder / "feedback" / ("feedback_" + digest(clean)[:20] + ".jsonl"), clean, jsonl=True)
        storage.put_feedback(ROOT,batch,clean,digest)
    return {"feedback_imported": len(clean)}


def export_review(sid, batch):
    folder = get_folder(sid, batch)
    proposals, feedback = state(folder)
    samples = storage.members(ROOT,batch)
    by_id = {x["sku"]: x for x in samples}
    fields = ["sku", "name_es_original", "name_es_normalized", *FIELDS, *EXTRA_FIELDS, "original_category_comparison_only", "source_reference", "source_row_ids", "snapshot_id", "rules_version", "annotation_hash", "human_review_status", "human_corrections", "human_review_note", "human_reviewer", "human_feedback_binding"]
    versions={x['rules_version'] for x in proposals.values()}
    version=next(iter(versions)) if len(versions)==1 else 'mixed_versions'
    output = ROOT / "reports" / sid / batch / version
    output.mkdir(parents=True, exist_ok=True)
    # Derived review view can be refreshed; raw data/proposals/feedback remain immutable.
    dest = output / "human_review.csv"
    with dest.with_suffix(".partial").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for sku, x in sorted(proposals.items()):
            fb = feedback.get(sku, {})
            binding='CURRENT' if fb and fb.get('expected_annotation_hash')==x['annotation_hash'] else ('STALE' if fb else 'NONE')
            if binding=='STALE':fb={}
            record = {k: x.get(k, "") for k in fields}
            record.update(original_category_comparison_only=by_id[sku]["category_original"], source_reference=by_id[sku]["source_reference"], source_row_ids=by_id[sku]["source_row_ids"], human_review_status=fb.get("review_status", "待审核"), human_corrections=fb.get("corrections", {}), human_review_note=fb.get("note", ""), human_reviewer=fb.get("reviewer", ""),human_feedback_binding=binding)
            writer.writerow({k: canonical(v) if isinstance(v, (dict, list)) else v for k, v in record.items()})
    os.replace(dest.with_suffix(".partial"), dest)
    ai_dest=output/"ai_learning_results.csv"
    ai_fields=["sku","name_es_original","name_es_normalized",*FIELDS,*EXTRA_FIELDS,"snapshot_id","source_row_id","rules_version","annotation_hash"]
    with ai_dest.with_suffix(".partial").open("w",encoding="utf-8-sig",newline="") as f:
        writer=csv.DictWriter(f,fieldnames=ai_fields);writer.writeheader()
        for sku,x in sorted(proposals.items()):
            writer.writerow({k:canonical(x[k]) if isinstance(x.get(k),(dict,list)) else x.get(k,'') for k in ai_fields})
    os.replace(ai_dest.with_suffix(".partial"),ai_dest)
    counts = Counter(x["review_status"] for x in proposals.values())
    types = Counter(x["category_l1"] for x in proposals.values())
    problems = [{"sku": sku, "name_es_original": x["name_es_original"], "issue": "信息不足" if x["review_status"] == "信息不足" else "LOW_OR_MEDIUM_CONFIDENCE", "evidence": x["evidence"]} for sku, x in proposals.items() if x["confidence"] != "高" or x["review_status"] == "信息不足"]
    immutable(output / ("issues_" + digest(problems)[:12] + ".json"), problems)
    result = {"snapshot_id": sid, "sample_id": batch, "processed": len(proposals), "unprocessed_in_full_copy": len(storage.load_products(ROOT,sid)) - len(proposals), "status_counts": dict(counts), "confidence_counts": dict(Counter(x["confidence"] for x in proposals.values())), "category_l1_counts": dict(types), "human_confirmed": sum(x.get("review_status") == "已确认" for x in feedback.values()), "issue_records": len(problems), "review_csv": str(dest), "ai_learning_csv":str(ai_dest), "rules_version": version}
    (output / "p1_summary.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def main():
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest="command", required=True)
    q = sub.add_parser("snapshot"); q.add_argument("--master", required=True)
    q = sub.add_parser("validate-snapshot"); q.add_argument("--snapshot", required=True)
    q = sub.add_parser("sample"); q.add_argument("--snapshot", required=True); q.add_argument("--seed", type=int, default=20261022)
    for name in ["import-analysis", "feedback", "export-review", "pending", "compare"]:
        q = sub.add_parser(name); q.add_argument("--snapshot", required=True); q.add_argument("--batch", required=True)
        if name in {"import-analysis", "feedback"}: q.add_argument("--input", required=True)
        if name == "import-analysis":
            q.add_argument("--reprocess", nargs="*", default=[])
            q.add_argument("--rules-version",default=RULE_VERSION)
    a = p.parse_args()
    if a.command == "snapshot": result = snapshot(a.master)
    elif a.command == "validate-snapshot": result = validate_snapshot(a.snapshot)
    elif a.command == "sample": result = sample(a.snapshot, a.seed)
    elif a.command == "import-analysis": result = import_analysis(a.snapshot, a.batch, a.input, a.reprocess,a.rules_version)
    elif a.command == "feedback": result = import_feedback(a.snapshot, a.batch, a.input)
    elif a.command == "export-review": result = export_review(a.snapshot, a.batch)
    elif a.command == "compare":
        folder = get_folder(a.snapshot, a.batch); history = defaultdict(list)
        for x in storage.history(ROOT,a.batch): history[x["sku"]].append({k: x.get(k) for k in ["rules_version", "annotation_hash", *FIELDS,*EXTRA_FIELDS]})
        result = {sku: v for sku, v in history.items() if len(v) > 1}
    else:
        folder = get_folder(a.snapshot, a.batch); proposals, feedback = state(folder)
        result = [x for x in read_jsonl(folder / "names_only_input.jsonl") if x["sku"] not in proposals and feedback.get(x["sku"], {}).get("review_status") != "已确认"]
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
