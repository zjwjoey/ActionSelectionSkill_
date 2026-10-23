"""Read historical Master title versions; skip held IDs before title access."""
import collections, hashlib, json, re, sys
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
import pipeline as p
from semantic22_prepare import OUT,BASE

def main():
    excluded=json.loads((OUT/'excluded_ids.json').read_text(encoding='utf8'))
    held={r['sku'] for r in excluded if r['state']=='RESERVED_HOLDOUT_NOT_USED'}
    # Sealed plan exposes only IDs and hashes, not title content.
    seal=json.loads((ROOT/'learning_data/p11/sealed_unseen50_plan.json').read_text(encoding='utf8'))
    sealed_hash={r['title_sha256'] for r in seal['reserved_ids']}
    current={r['sku']:r for r in p.read_jsonl(OUT/'inputs.jsonl')}
    excluded_ids={r['sku'] for r in excluded}
    files=[Path('F:/Action_Master/Action_Master.xlsx'),*sorted(Path('F:/ActionSKUTracker/runtime/master').glob('Action_Master_before*.xlsx'))]
    manifests=[];versions=collections.defaultdict(list);counts=collections.Counter();seen_hash=set()
    # Existing older source snapshot joins the same source-version provenance.
    for path in files:
        sha=p.filehash(path); info={'path':str(path),'sha256':sha}
        if sha in seen_hash:
            manifests.append(info|{'duplicate_file_content':True});continue
        seen_hash.add(sha);wb=load_workbook(path,read_only=True,data_only=True);local=collections.Counter()
        for sn,hr,sku_col in [('02_SKU_ES_CURRENT',1,'SKU'),('08_LONG_TERM_MASTER',7,'正式SKU')]:
            if sn not in wb:continue
            it=wb[sn].iter_rows(values_only=True)
            for _ in range(hr):head=next(it)
            index={s:i for i,s in enumerate(head) if s}
            for rownum,row in enumerate(it,hr+1):
                sku=p.sku_text(row[index[sku_col]])
                if not sku or not re.fullmatch(r'\d+',sku):local['unconfirmed_identity']+=1;continue
                if sku in held:local['sealed_sku_skipped_before_title']+=1;continue
                title=p.normalize_name(row[index['西班牙语品名']])
                if hashlib.sha256(title.encode()).hexdigest() in sealed_hash:local['sealed_title_hash_skipped']+=1;continue
                if not title or re.search('[\u3400-\u9fff\ufffd]',title):local['source_language_or_empty']+=1;continue
                ref={'file':str(path),'file_sha256':sha,'sheet':sn,'excel_row':rownum,'source_date':str(row[index['最后确认存在日期' if sn=='02_SKU_ES_CURRENT' else '最后观察日期']]) if ('最后确认存在日期' if sn=='02_SKU_ES_CURRENT' else '最后观察日期') in index else None}
                versions[(sku,title)].append(ref);local['valid_source_rows']+=1
        wb.close();assert p.filehash(path)==sha
        manifests.append(info|{'counts':dict(local)});counts.update(local)
    old=ROOT/'source_snapshot/master_d1af859e842adbe1/sku_source_rows.jsonl'
    for line in old.open(encoding='utf8'):
        r=json.loads(line);sku=r['sku']
        if sku in held:continue
        title=p.normalize_name(r['name_es_original'])
        if not title or re.search('[\u3400-\u9fff\ufffd]',title):continue
        if hashlib.sha256(title.encode()).hexdigest() in sealed_hash:continue
        versions[(sku,title)].append({'file':str(old),'source_row_id':r['source_row_id'],'snapshot_id':r['snapshot_id'],'source_date':r['source_date']})
    rows=[]
    for (sku,title),refs in sorted(versions.items()):
        state='SAME_CURRENT_TITLE' if sku in current and current[sku]['name_es_original']==title else 'HISTORICAL_TITLE_VARIANT' if sku in current else 'CURRENT_SOURCE_QUARANTINE' if sku in excluded_ids else 'EXTRA_HISTORICAL_SKU'
        rows.append({'sku':sku,'name_es_original':title,'source_title_hash':p.digest(title),'version_record_id':'hist_'+p.digest([sku,title])[:24],'state':state,'sources':refs})
    p.immutable(OUT/'historical_title_versions.jsonl',rows,jsonl=True)
    report={'historical_files':manifests,'source_version_records':len(rows),'distinct_skus':len({r['sku'] for r in rows}),'version_states':dict(collections.Counter(r['state'] for r in rows)),'additional_distinct_skus':len({r['sku'] for r in rows if r['state']=='EXTRA_HISTORICAL_SKU'}),'current_master_includes_long_term_skus':10121,'current_sheet_skus':5763,'long_term_not_in_current_sheet':4358,'archive_rows_without_confirmed_sku_not_promoted':6893,'sealed_title_payloads_exported':0}
    p.immutable(OUT/'historical_source_report.json',report);print(json.dumps(report,ensure_ascii=False))
if __name__=='__main__':main()
