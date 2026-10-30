"""Atomic, idempotent Excel review import; never overwrites AI records."""
import json
from pathlib import Path
from datetime import datetime,timezone
from runtime.v3.core import digest
from runtime.v3.workflow import filehash
from runtime.v31.state import State
KINDS=['接受匹配结果','修改最佳参考SKU','商品类型判断错误','商品主体判断错误','商品关系判断错误','规格判断错误','匹配分不合理','无匹配但存在正确参考','价格证据错误']

def import_feedback(out,path,reviewer):
 import openpyxl
 if not reviewer or not reviewer.strip():raise ValueError('REAL_REVIEWER_REQUIRED')
 out=Path(out);results={r['record_id']:r for r in map(json.loads,(out/'results.jsonl').read_text(encoding='utf8').splitlines())}
 w=openpyxl.load_workbook(path,data_only=False,read_only=True)
 try:
  if '商品匹配结果' not in w.sheetnames:raise ValueError('MATCH10_SHEET_REQUIRED')
  rows=list(w['商品匹配结果'].iter_rows())
 finally:w.close()
 head=[str(c.value or '') for c in rows[4]];events=[]
 required=['记录ID','结果hash','方法版本','来源hash','原始商品ID','原始商品名称','参考SKU','商品匹配分（10分制）','人工反馈类型','新参考SKU','人工意见','人工关系','建议分数']
 if any(k not in head for k in required):raise ValueError('REVIEW_TEMPLATE_COLUMNS_MISSING')
 indexes={k:head.index(k) for k in required}
 for rowno,cells in enumerate(rows[5:],6):
  def get(k):return cells[indexes[k]].value
  kind=get('人工反馈类型');note=get('人工意见')
  if not kind and not note and not get('新参考SKU') and not get('人工关系') and get('建议分数') is None:continue
  if any(c.data_type=='f' for c in cells):raise ValueError('REVIEW_FORMULAS_NOT_ALLOWED:'+str(rowno))
  rid=str(get('记录ID') or '');r=results.get(rid)
  if not r or get('结果hash')!=digest(r) or get('方法版本')!=r['method'] or get('来源hash')!=r['source']['source_sha256']:raise ValueError('STALE_OR_UNBOUND_FEEDBACK:'+str(rowno))
  if str(get('原始商品ID'))!=r['source']['source_product_id'] or get('原始商品名称')!=r['source']['title_original'] or str(get('参考SKU') or '')!=str((r['primary'] or {}).get('reference_id') or '') or get('商品匹配分（10分制）')!=r['match_score_10']:raise ValueError('READONLY_RESULT_CELLS_CHANGED')
  if kind not in KINDS or not isinstance(note,str) or not note.strip():raise ValueError('FEEDBACK_KIND_AND_REASON_REQUIRED')
  newsku=str(get('新参考SKU') or '').strip();newscore=get('建议分数');rel=get('人工关系')
  if kind in ['修改最佳参考SKU','无匹配但存在正确参考'] and not newsku:raise ValueError('NEW_REFERENCE_REQUIRED')
  if newscore is not None and (isinstance(newscore,bool) or not isinstance(newscore,(int,float)) or not 0<=newscore<=10 or round(newscore,1)!=newscore):raise ValueError('SCORE_0_10_ONE_DECIMAL_REQUIRED')
  if kind=='匹配分不合理' and (not newsku and not r['primary']):raise ValueError('SCORE_FEEDBACK_NEEDS_REFERENCE_AND_REASON')
  if rel and rel not in ['IDENTICAL','SAME_TYPE','FUNCTIONAL_ALTERNATIVE','SHARED_NEED','UNRELATED','INSUFFICIENT']:raise ValueError('REVIEW_RELATION_INVALID')
  events.append(dict(record_id=rid,result_hash=digest(r),method=r['method'],source_sha=r['source']['source_sha256'],original_reference=(r['primary'] or {}).get('reference_id'),original_analysis=r['analysis'],original_score=r['match_score_10'],reviewer=reviewer.strip(),kind=kind,new_reference=newsku or None,reason=note.strip(),new_relation=rel,new_score=newscore))
 w.close()
 if len({e['record_id'] for e in events})!=len(events):raise ValueError('DUPLICATE_PRODUCT_FEEDBACK_IN_IMPORT')
 # All rows validated before any writes. Dedup ignores file metadata and import time.
 state=State(out/'state.sqlite3');inserted=0
 try:
  existing={digest({k:v for k,v in e.items() if k not in ['reviewed_at','workbook_sha']}):e for _,e in state.all('match10_feedback')}
  state.db.execute('BEGIN')
  for e in events:
   if digest(e) in existing:continue
   payload={**e,'reviewed_at':datetime.now(timezone.utc).isoformat(),'workbook_sha':filehash(Path(path))}
   key='match10_feedback-'+digest(payload)[:24]
   from runtime.v3.core import canonical
   state.db.execute('INSERT INTO v31_events VALUES(?,?,?)',(key,'match10_feedback',canonical(payload)));inserted+=1
  state.db.commit();return dict(validated=len(events),inserted=inserted,duplicates=len(events)-inserted)
 except BaseException:state.db.rollback();raise
 finally:state.close()
